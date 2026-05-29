#!/usr/bin/env python3
"""
scanner_opcoes_supertrend.py
────────────────────────────────────────────────────────────────────────
Scanner automático de oportunidades em opções — SuperTrend v1

Uso:
    python -X utf8 scanner_opcoes_supertrend.py --patrimonio 10000 --salvar
    python -X utf8 scanner_opcoes_supertrend.py --sem-iv
    python -X utf8 scanner_opcoes_supertrend.py --ativos PETR4.SA BBAS3.SA

AVISO: Material exclusivamente educacional. Não é recomendação de
investimento. Operações com opções envolvem risco de perda total.
────────────────────────────────────────────────────────────────────────
"""

import argparse
import sys
import webbrowser
from datetime import datetime, date, timedelta
from pathlib import Path

import numpy as np
import pandas as pd
import yfinance as yf

if sys.stdout.encoding and sys.stdout.encoding.lower() != "utf-8":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
        sys.stderr.reconfigure(encoding="utf-8")
    except AttributeError:
        pass

try:
    import iv_b3
    _IV_DISPONIVEL = True
except ImportError:
    _IV_DISPONIVEL = False


# ══════════════════════════════════════════════════════════════════════
# PARÂMETROS
# ══════════════════════════════════════════════════════════════════════

ATIVOS_PADRAO = [
    "PETR4.SA", "VALE3.SA", "CSNA3.SA", "USIM5.SA",
    "GOAU4.SA", "GGBR4.SA", "MGLU3.SA", "BRAV3.SA",
    "BRAP4.SA", "CMIN3.SA", "ASAI3.SA",
]

ATR_PERIODO     = 10
ATR_MULT_PADRAO = 3.0
JANELA_LATERAL  = 10
MAX_VIRADAS     = 3
VOL_THRESHOLD   = 0.35
DU_MIN          = 20
DU_MAX          = 40
REGRA_1PCT      = 0.01

CATALOGO: dict[str, list[tuple[str, str]]] = {
    "alta_vol_baixa": [
        ("1 – Compra a Seco de Call",   "Iniciante"),
        ("3 – Trava de Alta c/ Call",   "Iniciante"),
        ("16 – Lançamento Coberto",     "Iniciante (precisa ter a ação)"),
    ],
    "alta_vol_alta": [
        ("2 – Venda a Seco de Put",     "Intermediário (exige margem)"),
        ("4 – Trava de Alta c/ Put",    "Intermediário"),
    ],
    "baixa_vol_baixa": [
        ("5 – Compra a Seco de Put",    "Iniciante"),
        ("7 – Trava de Baixa c/ Put",   "Intermediário"),
    ],
    "baixa_vol_alta": [
        ("6 – Venda a Seco de Call",    "Avançado (exige cobertura)"),
        ("8 – Trava de Baixa c/ Call",  "Intermediário"),
    ],
    "lateral": [
        ("9 – Trava de Linha",              "Avançado"),
        ("10 – Straddle Vendido",           "Avançado (exige margem)"),
        ("11 – Straddle Sint. Vendido",     "Avançado (exige a ação)"),
    ],
    "explosao": [
        ("12 – Straddle Comprado",          "Intermediário"),
        ("13 – Straddle Sint. Comprado",    "Intermediário"),
    ],
}


# ══════════════════════════════════════════════════════════════════════
# SUPERTREND
# ══════════════════════════════════════════════════════════════════════

def _atr_rma(df: pd.DataFrame, periodo: int) -> pd.Series:
    h, l, c = df["High"], df["Low"], df["Close"]
    tr = pd.concat(
        [h - l, (h - c.shift()).abs(), (l - c.shift()).abs()],
        axis=1,
    ).max(axis=1)
    return tr.ewm(alpha=1.0 / periodo, adjust=False).mean()


def calcular_supertrend(df, periodo=ATR_PERIODO, mult=ATR_MULT_PADRAO):
    hl2    = (df["High"] + df["Low"]) / 2
    atr    = _atr_rma(df, periodo)
    ub_raw = (hl2 + mult * atr).values
    lb_raw = (hl2 - mult * atr).values
    close  = df["Close"].values
    n      = len(df)

    final_upper = ub_raw.copy()
    final_lower = lb_raw.copy()
    for i in range(1, n):
        final_upper[i] = (
            ub_raw[i]
            if ub_raw[i] < final_upper[i - 1] or close[i - 1] > final_upper[i - 1]
            else final_upper[i - 1]
        )
        final_lower[i] = (
            lb_raw[i]
            if lb_raw[i] > final_lower[i - 1] or close[i - 1] < final_lower[i - 1]
            else final_lower[i - 1]
        )

    st   = np.empty(n)
    dire = np.empty(n, dtype=int)
    st[0]   = final_upper[0]
    dire[0] = -1

    for i in range(1, n):
        if st[i - 1] == final_upper[i - 1]:
            if close[i] > final_upper[i]:
                st[i], dire[i] = final_lower[i], 1
            else:
                st[i], dire[i] = final_upper[i], -1
        else:
            if close[i] < final_lower[i]:
                st[i], dire[i] = final_upper[i], -1
            else:
                st[i], dire[i] = final_lower[i], 1

    return pd.DataFrame({"ST_valor": st, "ST_direcao": dire}, index=df.index)


def detectar_lateral(direcao, janela, max_viradas):
    recente = direcao.iloc[-janela:]
    return int((recente.diff().abs() > 0).sum()) > max_viradas


def vol_historica_anual(df, janela=20):
    ret = np.log(df["Close"] / df["Close"].shift()).dropna()
    return float(ret.tail(janela).std() * np.sqrt(252))


# ══════════════════════════════════════════════════════════════════════
# VENCIMENTOS B3
# ══════════════════════════════════════════════════════════════════════

def vencimentos_b3(n_meses=3):
    hoje = date.today()
    resultado = []
    for delta in range(n_meses + 1):
        mes = (hoje.month - 1 + delta) % 12 + 1
        ano = hoje.year + (hoje.month - 1 + delta) // 12
        primeiro = date(ano, mes, 1)
        sextas = sorted(
            d for d in (primeiro + timedelta(days=i) for i in range(31))
            if d.month == mes and d.weekday() == 4
        )
        if len(sextas) >= 3:
            resultado.append(sextas[2])
    return sorted(set(resultado))


def dias_uteis_entre(inicio, fim):
    total, d = 0, inicio
    while d < fim:
        if d.weekday() < 5:
            total += 1
        d += timedelta(days=1)
    return total


def melhor_vencimento(vencimentos):
    hoje = date.today()
    for v in vencimentos:
        du = dias_uteis_entre(hoje, v)
        if DU_MIN <= du <= DU_MAX:
            return v, du
    for v in vencimentos:
        du = dias_uteis_entre(hoje, v)
        if du > 0:
            return v, du
    return None, 0


# ══════════════════════════════════════════════════════════════════════
# ANÁLISE POR ATIVO
# ══════════════════════════════════════════════════════════════════════

def analisar(ticker, mult=ATR_MULT_PADRAO):
    base = ticker.replace(".SA", "")
    resultado = {"ticker": base, "ok": False}

    try:
        raw = yf.download(
            ticker, period="6mo", interval="1d",
            progress=False, auto_adjust=True
        )
        if isinstance(raw.columns, pd.MultiIndex):
            raw.columns = raw.columns.get_level_values(0)

        if len(raw) < ATR_PERIODO + 5:
            resultado["erro"] = "Dados insuficientes"
            return resultado

        st_df = calcular_supertrend(raw, ATR_PERIODO, mult)
        dados = pd.concat([raw, st_df], axis=1).dropna(subset=["ST_valor"])

        preco   = float(dados["Close"].iloc[-1])
        st_val  = float(dados["ST_valor"].iloc[-1])
        direcao = int(dados["ST_direcao"].iloc[-1])
        vol_h   = vol_historica_anual(dados, 20)
        lateral = detectar_lateral(dados["ST_direcao"], JANELA_LATERAL, MAX_VIRADAS)

        venc_list = vencimentos_b3()
        venc, du  = melhor_vencimento(venc_list)
        du_ok     = (DU_MIN <= du <= DU_MAX) if venc else False

        vol_str = "alta" if vol_h >= VOL_THRESHOLD else "baixa"
        if lateral:
            dir_str = "LATERAL"
            chave   = "lateral"
        elif direcao == 1:
            dir_str = "ALTA"
            chave   = f"alta_vol_{vol_str}"
        else:
            dir_str = "BAIXA"
            chave   = f"baixa_vol_{vol_str}"

        resultado.update(
            ok=True,
            preco=preco,
            st_valor=st_val,
            dist_pct=(preco - st_val) / st_val * 100,
            direcao_num=direcao,
            direcao_str=dir_str,
            lateral=lateral,
            vol_hist=vol_h,
            vol_decisao=vol_h,
            vol_str=vol_str,
            vol_fonte="Vol hist. 20d (proxy)",
            iv_real=None,
            chave=chave,
            estrategias=CATALOGO.get(chave, []),
            vencimento=venc,
            du=du,
            du_ok=du_ok,
            estrategias_detalhadas=[],
        )

    except Exception as exc:
        resultado["erro"] = str(exc)

    return resultado


def _aplicar_iv_real(resultado, iv_real):
    if not resultado.get("ok") or iv_real is None:
        return resultado

    r = resultado.copy()
    r["iv_real"]     = iv_real
    r["vol_decisao"] = iv_real
    r["vol_fonte"]   = "IV real B3 (COTAHIST)"

    vol_str_novo = "alta" if iv_real >= VOL_THRESHOLD else "baixa"
    r["vol_str"] = vol_str_novo

    if not r["lateral"]:
        dir_prefix   = "alta" if r["direcao_num"] == 1 else "baixa"
        chave_novo   = f"{dir_prefix}_vol_{vol_str_novo}"
        r["chave"]   = chave_novo
        r["estrategias"] = CATALOGO.get(chave_novo, [])

    return r


def _enriquecer_legs(resultado, opcoes_df, cdi):
    if not resultado.get("ok") or opcoes_df is None or opcoes_df.empty:
        return resultado

    r = resultado.copy()
    try:
        detalhes = iv_b3.selecionar_detalhes_estrategias(
            r["ticker"], r["preco"], opcoes_df,
            r["chave"], DU_MIN, DU_MAX,
        )
        r["estrategias_detalhadas"] = detalhes
    except Exception:
        r["estrategias_detalhadas"] = []
    return r


def _montar_checklist(r):
    return {
        "Tendência definida (ST estável)":      not r["lateral"],
        "Vencimento no prazo ideal (20–40 du)": r["du_ok"],
        "Vol fonte confiável":                  r["vol_fonte"] != "Vol hist. 20d (proxy)",
        "Strike c/ delta razoável (verificar)": True,
        "Alvo 3:1 identificável (verificar)":   True,
    }


# ══════════════════════════════════════════════════════════════════════
# RELATÓRIO TEXTO
# ══════════════════════════════════════════════════════════════════════

_SEP  = "─" * 72
_SEP2 = "═" * 72


def formatar_relatorio(resultados, patrimonio):
    hoje = datetime.now().strftime("%d/%m/%Y %H:%M")
    linhas = [
        _SEP2,
        "  SCANNER DE OPÇÕES — SUPERTREND v1  (com IV real B3)",
        f"  Gerado em: {hoje}",
        f"  SuperTrend: ATR {ATR_PERIODO} | Mult. {ATR_MULT_PADRAO} "
        f"| Vol. threshold: {VOL_THRESHOLD*100:.0f}% a.a.",
        _SEP2,
    ]

    if patrimonio:
        risco  = patrimonio * REGRA_1PCT
        ex_min = patrimonio * 0.05
        ex_max = patrimonio * 0.20
        linhas += [
            "",
            f"  Patrimônio declarado : R$ {patrimonio:>12,.2f}",
            f"  Risco máx/op  (1%)   : R$ {risco:>12,.2f}",
            f"  Exposição em opções  : R$ {ex_min:,.2f}  –  R$ {ex_max:,.2f}  (5–20%)",
        ]

    sinais = [r for r in resultados if r.get("ok")]
    erros  = [r for r in resultados if not r.get("ok")]

    linhas += ["", _SEP, f"  SINAIS — {len(sinais)} ativo(s)", _SEP]

    for r in sinais:
        checks   = _montar_checklist(r)
        n_ok     = sum(checks.values())
        dir_ico  = "▲" if r["direcao_str"] == "ALTA" else ("▼" if r["direcao_str"] == "BAIXA" else "↔")
        venc_str = r["vencimento"].strftime("%d/%m/%Y") if r["vencimento"] else "N/A"
        du_tag   = f"{r['du']} du" + ("" if r["du_ok"] else "  ⚠ fora da janela")

        if r["iv_real"] is not None:
            vol_linha = f"Vol hist: {r['vol_hist']*100:.1f}%  |  IV real B3: {r['iv_real']*100:.1f}%  →  {r['vol_str'].upper()}"
        else:
            vol_linha = f"Vol hist: {r['vol_hist']*100:.1f}%  →  {r['vol_str'].upper()}  ⚠ IV real indisponível"

        linhas += [
            "",
            f"  {dir_ico}  {r['ticker']:<10}  Preço: R$ {r['preco']:>8.2f}  |  ST: R$ {r['st_valor']:>8.2f}  ({r['dist_pct']:+.1f}%)",
            f"     Direção: {r['direcao_str']:<8}  |  {vol_linha}",
            f"     Vencimento sugerido: {venc_str}  ({du_tag})",
        ]

        for det in r.get("estrategias_detalhadas", []):
            linhas.append(f"     ► {det['nome']}  [{det['nivel']}]")
            if det["legs"]:
                for lg in det["legs"]:
                    linhas.append(
                        f"         {lg['acao']:<8}  {lg['codneg']:<12}"
                        f"  Strike R${lg['strike']:.2f}  Prêmio R${lg['preult']:.2f}"
                        f"  Venc {lg['vencimento'].strftime('%d/%m/%Y')}  ({lg['du']} du)"
                    )
                po = det.get("preco_op", {})
                if po.get("entrada") is not None:
                    stop_str = f"R${po['stop']:.2f}" if po.get("stop") else "—"
                    alvo_str = f"R${po['alvo']:.2f}" if po.get("alvo") else "—"
                    linhas.append(f"         Entrada: R${po['entrada']:.2f}  Stop: {stop_str}  Alvo: {alvo_str}  R/R: {po.get('rr','')}")
                if po.get("descricao"):
                    linhas.append(f"         {po['descricao']}")
            else:
                linhas.append(f"         (opções não encontradas no COTAHIST para este vencimento)")

        linhas.append(f"     Checklist ({n_ok}/{len(checks)} OK):")
        for item, ok in checks.items():
            linhas.append(f"       {'✓' if ok else '⚠'} {item}")

    if erros:
        linhas += ["", _SEP, f"  ERROS — {len(erros)} ativo(s)", _SEP]
        for r in erros:
            linhas.append(f"  ✗  {r['ticker']:<10}  {r.get('erro', 'erro desconhecido')}")

    linhas += [
        "", _SEP,
        "  MAPA DE DECISÃO RÁPIDO", _SEP,
        "  ACIMA do ST + vol BAIXA  →  Estratégias 1, 3, 16",
        "  ACIMA do ST + vol ALTA   →  Estratégias 2, 4",
        "  ABAIXO do ST + vol BAIXA →  Estratégias 5, 7",
        "  ABAIXO do ST + vol ALTA  →  Estratégias 6, 8",
        "  LATERAL (ST serrando)    →  Estratégias 9, 10, 11",
        "  Evento marcado           →  Estratégias 12, 13",
        _SEP,
        "  • Aja SEMPRE no fechamento do candle diário.",
        "  • 1% do patrimônio por operação; 5–20% de exposição total.",
        "  • Vencimento ideal: 20–40 dias úteis.",
        "  • Material educacional — não é recomendação de investimento.",
        _SEP2,
    ]

    return "\n".join(linhas)


# ══════════════════════════════════════════════════════════════════════
# EXPORTAR EXCEL
# ══════════════════════════════════════════════════════════════════════

def exportar_excel(resultados, caminho):
    rows = []
    for r in resultados:
        checks = _montar_checklist(r) if r.get("ok") else {}
        if not r.get("ok"):
            rows.append({
                "Ativo": r["ticker"], "Preço": None, "SuperTrend": None,
                "Direção": "ERRO", "Dist. % ST": None,
                "Vol Hist. (%)": None, "IV Real B3 (%)": None,
                "Vol Usada (%)": None, "Fonte Vol": None,
                "Nível Vol": None, "Cenário": None,
                "Estratégias": r.get("erro", ""),
                "Opção 1": None, "Ação 1": None, "Prêmio 1": None,
                "Opção 2": None, "Ação 2": None, "Prêmio 2": None,
                "Entrada op.": None, "Stop": None, "Alvo": None,
                "Vencimento": None, "Dias Úteis": None, "Checklist OK": None,
            })
        else:
            strats = " / ".join(n for n, _ in r["estrategias"])
            dets   = r.get("estrategias_detalhadas", [])

            row = {
                "Ativo":          r["ticker"],
                "Preço":          round(r["preco"], 2),
                "SuperTrend":     round(r["st_valor"], 2),
                "Direção":        r["direcao_str"],
                "Dist. % ST":     round(r["dist_pct"], 2),
                "Vol Hist. (%)":  round(r["vol_hist"] * 100, 1),
                "IV Real B3 (%)": round(r["iv_real"] * 100, 1) if r["iv_real"] else "N/D",
                "Vol Usada (%)":  round(r["vol_decisao"] * 100, 1),
                "Fonte Vol":      r["vol_fonte"],
                "Nível Vol":      r["vol_str"].upper(),
                "Cenário":        r["chave"].replace("_", " ").upper(),
                "Estratégias":    strats,
                "Vencimento":     r["vencimento"].strftime("%d/%m/%Y") if r["vencimento"] else "N/A",
                "Dias Úteis":     r["du"],
                "Checklist OK":   f"{sum(checks.values())}/{len(checks)}",
            }
            for d in dets:
                num = d["nome"].split("–")[0].strip()
                legs = d["legs"]
                po   = d.get("preco_op", {})
                row[f"[{num}] Leg1"]    = legs[0]["codneg"] if len(legs) > 0 else ""
                row[f"[{num}] Ação1"]   = legs[0]["acao"]   if len(legs) > 0 else ""
                row[f"[{num}] Leg2"]    = legs[1]["codneg"] if len(legs) > 1 else ""
                row[f"[{num}] Ação2"]   = legs[1]["acao"]   if len(legs) > 1 else ""
                row[f"[{num}] Entrada"] = po.get("entrada", "")
                row[f"[{num}] Stop"]    = po.get("stop", "")
                row[f"[{num}] Alvo"]    = po.get("alvo", "")
            rows.append(row)

    df = pd.DataFrame(rows)

    with pd.ExcelWriter(caminho, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Scanner")
        ws = writer.sheets["Scanner"]

        for col in ws.columns:
            max_len = max((len(str(c.value or "")) for c in col), default=10) + 4
            ws.column_dimensions[col[0].column_letter].width = min(max_len, 40)

        from openpyxl.styles import PatternFill
        cores = {"ALTA": "D6F5D6", "BAIXA": "FAD7D7", "LATERAL": "FFF4CC", "ERRO": "EEEEEE"}
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
            dir_val = str(row[3].value or "")
            cor  = cores.get(dir_val, "FFFFFF")
            fill = PatternFill(start_color=cor, end_color=cor, fill_type="solid")
            for cell in row:
                cell.fill = fill

    print(f"  Excel salvo em: {caminho}")


# ══════════════════════════════════════════════════════════════════════
# RELATÓRIO HTML
# ══════════════════════════════════════════════════════════════════════

def _badge_dir(d):
    if d == "ALTA":
        return '<span class="badge alta">▲ ALTA</span>'
    if d == "BAIXA":
        return '<span class="badge baixa">▼ BAIXA</span>'
    return '<span class="badge lateral">↔ LATERAL</span>'


def _badge_vol(v):
    cls = "vol-alta" if v == "alta" else "vol-baixa"
    txt = "Vol ALTA" if v == "alta" else "Vol BAIXA"
    return f'<span class="badge {cls}">{txt}</span>'


def _nivel_badge(nivel):
    cls = "nivel-ini" if "Iniciante" in nivel else ("nivel-int" if "Intermediário" in nivel else "nivel-ava")
    return f'<span class="nivel {cls}">{nivel}</span>'


def gerar_html(resultados, patrimonio, timestamp):
    sinais = [r for r in resultados if r.get("ok")]
    erros  = [r for r in resultados if not r.get("ok")]

    cards_html = ""
    for r in sinais:
        checks = _montar_checklist(r)
        n_ok   = sum(checks.values())

        iv_linha = (
            f"<span class='iv-real'>IV real B3: <b>{r['iv_real']*100:.1f}%</b></span>"
            if r["iv_real"] else
            f"<span class='iv-hist'>Vol hist: <b>{r['vol_hist']*100:.1f}%</b> ⚠ sem IV real</span>"
        )

        dets = r.get("estrategias_detalhadas", [])
        strats_legs_html = ""
        for det in dets:
            nome_det  = det["nome"]
            nivel_det = det["nivel"]
            legs      = det["legs"]
            po        = det.get("preco_op", {})

            # Tabela de legs
            if legs:
                tbl = "<table class='legs-table'><thead><tr><th>Ação</th><th>Ticker</th><th>Tipo</th><th>Strike</th><th>Prêmio</th><th>Vencimento</th><th>du</th></tr></thead><tbody>"
                for lg in legs:
                    acao_cls = "comprar" if lg["acao"] == "COMPRAR" else "vender"
                    tbl += (
                        f"<tr>"
                        f"<td><span class='acao {acao_cls}'>{lg['acao']}</span></td>"
                        f"<td><b>{lg['codneg']}</b></td>"
                        f"<td>{lg['tipo']}</td>"
                        f"<td>R$ {lg['strike']:.2f}</td>"
                        f"<td><b>R$ {lg['preult']:.2f}</b></td>"
                        f"<td>{lg['vencimento'].strftime('%d/%m/%Y')}</td>"
                        f"<td>{lg['du']}</td>"
                        f"</tr>"
                    )
                tbl += "</tbody></table>"

                entrada_val = po.get("entrada")
                stop_val    = po.get("stop")
                alvo_val    = po.get("alvo")
                preco_bloco = f"""
                <div class='preco-op'>
                  <div class='po-item entrada'>Entrada<br><b>{'R$ ' + f"{entrada_val:.2f}" if entrada_val is not None else '—'}</b></div>
                  <div class='po-item stop'>Stop<br><b>{'R$ ' + f"{stop_val:.2f}" if stop_val is not None else '—'}</b></div>
                  <div class='po-item alvo'>Alvo<br><b>{'R$ ' + f"{alvo_val:.2f}" if alvo_val is not None else '—'}</b></div>
                  <div class='po-item rr'>R / R<br><b>{po.get('rr','—')}</b></div>
                </div>
                <p class='po-desc'>{po.get('descricao','')}</p>"""
            else:
                tbl = ""
                preco_bloco = f"<p class='sem-legs'>{po.get('descricao', 'Opções não encontradas no COTAHIST para este vencimento.')}</p>"

            strats_legs_html += f"""
            <div class='estrategia-bloco'>
              <div class='estrategia-titulo'>{nome_det} &nbsp; {_nivel_badge(nivel_det)}</div>
              {tbl}
              {preco_bloco}
            </div>"""

        checklist_html = ""
        for item, ok in checks.items():
            ico = "✓" if ok else "⚠"
            cls = "check-ok" if ok else "check-warn"
            checklist_html += f"<li class='{cls}'>{ico} {item}</li>"

        venc_str = r["vencimento"].strftime("%d/%m/%Y") if r["vencimento"] else "N/A"
        du_warn  = "" if r["du_ok"] else " ⚠"

        dir_cls = r["direcao_str"].lower()
        cards_html += f"""
        <div class="card card-{dir_cls}">
          <div class="card-header">
            <div class="ticker-line">
              <span class="ticker">{r['ticker']}</span>
              {_badge_dir(r['direcao_str'])}
              {_badge_vol(r['vol_str'])}
              <span class="checklist-badge">{n_ok}/{len(checks)} ✓</span>
            </div>
            <div class="preco-line">
              <span>Preço: <b>R$ {r['preco']:.2f}</b></span>
              <span>SuperTrend: <b>R$ {r['st_valor']:.2f}</b></span>
              <span class="dist">Distância: {r['dist_pct']:+.1f}%</span>
            </div>
            <div class="vol-line">
              {iv_linha}
              &nbsp;|&nbsp; Vol hist: {r['vol_hist']*100:.1f}%
              &nbsp;|&nbsp; Venc: <b>{venc_str}</b> ({r['du']} du{du_warn})
            </div>
          </div>

          <div class="section-title">Estratégias e opções sugeridas (COTAHIST B3)</div>
          {strats_legs_html}

          <div class="section-title">Checklist</div>
          <ul class="checklist">{checklist_html}</ul>
        </div>"""

    erros_html = ""
    if erros:
        erros_html = "<div class='erros'><h3>Sem dados</h3><ul>"
        for r in erros:
            erros_html += f"<li>{r['ticker']} — {r.get('erro','erro desconhecido')}</li>"
        erros_html += "</ul></div>"

    pat_html = ""
    if patrimonio:
        pat_html = f"""
        <div class="pat-bar">
          Patrimônio: <b>R$ {patrimonio:,.2f}</b> &nbsp;|&nbsp;
          Risco máx/op (1%): <b>R$ {patrimonio*0.01:,.2f}</b> &nbsp;|&nbsp;
          Exposição total (5–20%): <b>R$ {patrimonio*0.05:,.2f} – R$ {patrimonio*0.20:,.2f}</b>
        </div>"""

    html = f"""<!DOCTYPE html>
<html lang="pt-BR">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>Scanner de Opções — SuperTrend v1</title>
<style>
  * {{ box-sizing: border-box; margin: 0; padding: 0; }}
  body {{ font-family: 'Segoe UI', Arial, sans-serif; background: #0f172a; color: #e2e8f0; min-height: 100vh; }}

  .header {{ background: linear-gradient(135deg, #1e3a5f, #0f2a45); padding: 28px 32px; border-bottom: 2px solid #334155; }}
  .header h1 {{ font-size: 1.6rem; color: #38bdf8; margin-bottom: 6px; }}
  .header p  {{ color: #94a3b8; font-size: 0.85rem; }}

  .pat-bar {{ background: #1e293b; padding: 10px 32px; font-size: 0.85rem; color: #94a3b8; border-bottom: 1px solid #334155; }}
  .pat-bar b {{ color: #e2e8f0; }}

  .disclaimer {{ background: #1c1510; color: #fb923c; text-align: center; padding: 8px; font-size: 0.78rem; border-bottom: 1px solid #92400e; }}

  .grid {{ display: grid; grid-template-columns: repeat(auto-fill, minmax(480px, 1fr)); gap: 20px; padding: 24px 32px; }}

  .card {{ background: #1e293b; border-radius: 12px; overflow: hidden; border: 1px solid #334155; }}
  .card-alta   {{ border-top: 4px solid #22c55e; }}
  .card-baixa  {{ border-top: 4px solid #ef4444; }}
  .card-lateral {{ border-top: 4px solid #f59e0b; }}

  .card-header {{ padding: 16px 20px 14px; background: #162032; }}
  .ticker-line {{ display: flex; align-items: center; gap: 10px; margin-bottom: 8px; flex-wrap: wrap; }}
  .ticker {{ font-size: 1.4rem; font-weight: 700; color: #f1f5f9; }}
  .preco-line  {{ font-size: 0.88rem; color: #94a3b8; margin-bottom: 4px; display: flex; gap: 16px; flex-wrap: wrap; }}
  .preco-line b {{ color: #e2e8f0; }}
  .dist {{ color: #64748b; }}
  .vol-line {{ font-size: 0.82rem; color: #64748b; }}

  .badge {{ padding: 3px 10px; border-radius: 20px; font-size: 0.78rem; font-weight: 600; }}
  .alta     {{ background: #14532d; color: #4ade80; }}
  .baixa    {{ background: #450a0a; color: #f87171; }}
  .lateral  {{ background: #451a03; color: #fbbf24; }}
  .vol-alta  {{ background: #450a0a; color: #fca5a5; }}
  .vol-baixa {{ background: #0c2340; color: #93c5fd; }}

  .checklist-badge {{ margin-left: auto; background: #334155; padding: 3px 10px; border-radius: 20px; font-size: 0.78rem; color: #94a3b8; }}

  .iv-real {{ color: #34d399; }} .iv-hist {{ color: #f59e0b; }}

  .nivel {{ padding: 2px 8px; border-radius: 12px; font-size: 0.72rem; font-weight: 600; }}
  .nivel-ini {{ background: #14532d; color: #4ade80; }}
  .nivel-int {{ background: #1e3a5f; color: #60a5fa; }}
  .nivel-ava {{ background: #4a044e; color: #e879f9; }}

  .section-title {{ font-size: 0.72rem; font-weight: 700; color: #64748b; text-transform: uppercase; letter-spacing: 0.08em; padding: 12px 20px 4px; }}

  .estrategia-bloco {{ border-top: 1px solid #1e293b; padding: 10px 0 4px; margin: 0 20px; }}
  .estrategia-bloco:first-child {{ border-top: none; }}
  .estrategia-titulo {{ font-size: 0.88rem; font-weight: 600; color: #e2e8f0; padding: 4px 0 8px; display: flex; align-items: center; gap: 8px; flex-wrap: wrap; }}

  .legs-table {{ width: 100%; border-collapse: collapse; font-size: 0.82rem; margin: 4px 0; }}
  .legs-table th {{ background: #0f172a; color: #64748b; font-weight: 600; padding: 6px 10px; text-align: left; font-size: 0.72rem; text-transform: uppercase; }}
  .legs-table td {{ padding: 7px 10px; border-bottom: 1px solid #1e293b; color: #cbd5e1; }}
  .legs-table tbody tr:hover {{ background: #162032; }}

  .acao {{ padding: 2px 8px; border-radius: 4px; font-size: 0.75rem; font-weight: 700; }}
  .comprar {{ background: #14532d; color: #4ade80; }}
  .vender  {{ background: #450a0a; color: #f87171; }}

  .preco-op {{ display: flex; gap: 10px; padding: 12px 20px 4px; flex-wrap: wrap; }}
  .po-item  {{ flex: 1; min-width: 80px; background: #0f172a; border-radius: 8px; padding: 8px 12px; text-align: center; font-size: 0.78rem; color: #94a3b8; }}
  .po-item b {{ display: block; font-size: 1rem; margin-top: 2px; }}
  .entrada b {{ color: #60a5fa; }}
  .stop    b {{ color: #f87171; }}
  .alvo    b {{ color: #4ade80; }}
  .rr      b {{ color: #fbbf24; }}

  .po-desc {{ font-size: 0.78rem; color: #64748b; padding: 4px 20px 12px; }}
  .sem-legs {{ font-size: 0.80rem; color: #475569; padding: 8px 20px 12px; font-style: italic; }}

  .checklist {{ list-style: none; padding: 4px 20px 16px; display: flex; flex-direction: column; gap: 5px; }}
  .checklist li {{ font-size: 0.82rem; }}
  .check-ok   {{ color: #4ade80; }}
  .check-warn {{ color: #f59e0b; }}

  .erros {{ background: #1e293b; margin: 0 32px 24px; border-radius: 10px; padding: 16px 20px; border-left: 4px solid #ef4444; }}
  .erros h3 {{ color: #f87171; margin-bottom: 8px; font-size: 0.9rem; }}
  .erros li {{ color: #94a3b8; font-size: 0.85rem; margin-top: 4px; list-style: none; }}

  .footer {{ text-align: center; color: #334155; font-size: 0.75rem; padding: 24px; border-top: 1px solid #1e293b; }}
</style>
</head>
<body>

<div class="header">
  <h1>Scanner de Opções — SuperTrend v1</h1>
  <p>Gerado em {timestamp} &nbsp;|&nbsp; ATR {ATR_PERIODO} &nbsp;|&nbsp; Mult. {ATR_MULT_PADRAO} &nbsp;|&nbsp; Vol threshold: {VOL_THRESHOLD*100:.0f}% a.a. &nbsp;|&nbsp; {len(sinais)} ativo(s) com sinal</p>
</div>

<div class="disclaimer">
  Material exclusivamente educacional. Não constitui recomendação de investimento. Operações com opções envolvem risco de perda total do capital investido.
</div>

{pat_html}

<div class="grid">
{cards_html}
</div>

{erros_html}

<div class="footer">
  Dados de mercado: B3 COTAHIST (gratuito) + yfinance + BACEN API &nbsp;|&nbsp;
  IV calculada via Black-Scholes + bisseção &nbsp;|&nbsp;
  SuperTrend idêntico ao TradingView (ATR Wilder's RMA)
</div>

</body>
</html>"""

    return html


# ══════════════════════════════════════════════════════════════════════
# MAIN
# ══════════════════════════════════════════════════════════════════════

def main():
    parser = argparse.ArgumentParser(description="Scanner de Opções — SuperTrend v1")
    parser.add_argument("--patrimonio", type=float, default=None)
    parser.add_argument("--mult",       type=float, default=ATR_MULT_PADRAO)
    parser.add_argument("--salvar",     action="store_true")
    parser.add_argument("--sem-iv",     action="store_true")
    parser.add_argument("--ativos",     nargs="+", default=None)
    parser.add_argument("--no-browser", action="store_true", help="Não abre o HTML automaticamente")
    args = parser.parse_args()

    ativos  = args.ativos or ATIVOS_PADRAO
    usar_iv = _IV_DISPONIVEL and not args.sem_iv

    print(f"\nScanner de Opções — SuperTrend v1")
    print(f"ATR {ATR_PERIODO} | Mult. {args.mult} | {len(ativos)} ativo(s)")
    if not _IV_DISPONIVEL:
        print("  [AVISO] iv_b3.py não encontrado — usando vol. histórica.")
    elif args.sem_iv:
        print("  [INFO] --sem-iv: IV real desativada.")
    else:
        print("  [INFO] IV real via COTAHIST B3 ativada.")
    print()

    # ── Etapa 1 ──────────────────────────────────────────────────────
    print("1/3  Baixando dados históricos...")
    resultados_base = {}
    for ticker in ativos:
        base = ticker.replace(".SA", "")
        print(f"  {base:<8}", end="  ", flush=True)
        r = analisar(ticker, mult=args.mult)
        resultados_base[ticker] = r
        if r.get("ok"):
            print(f"{r['direcao_str']:<8}  Vol hist: {r['vol_hist']*100:.1f}%  ST: R${r['st_valor']:.2f}  ✓")
        else:
            print(f"✗  {r.get('erro', 'erro')}")

    # ── Etapa 2 ──────────────────────────────────────────────────────
    iv_reais   = {}
    opcoes_df  = None
    cdi        = iv_b3.CDI_PADRAO if _IV_DISPONIVEL else 0.1325

    if usar_iv:
        print("\n2/3  Buscando IV real no COTAHIST B3...")
        precos_spot = {
            t.replace(".SA", ""): r["preco"]
            for t, r in resultados_base.items() if r.get("ok")
        }
        try:
            resultado_iv = iv_b3.obter_iv_todos(
                list(precos_spot.keys()), precos_spot,
                verbose=True, retornar_opcoes=True,
            )
            iv_reais, opcoes_df, cdi = resultado_iv
            n_iv = sum(1 for v in iv_reais.values() if v is not None)
            print(f"  → IV calculada para {n_iv}/{len(precos_spot)} ativos")
        except Exception as exc:
            print(f"  [AVISO] Falha no COTAHIST: {exc}")
    else:
        print("\n2/3  (IV real desativada)")

    # ── Etapa 3 ──────────────────────────────────────────────────────
    print("\n3/3  Consolidando e selecionando opções...")
    resultados = []
    for ticker in ativos:
        base = ticker.replace(".SA", "")
        r    = resultados_base[ticker]
        r    = _aplicar_iv_real(r, iv_reais.get(base))
        r    = _enriquecer_legs(r, opcoes_df, cdi)
        resultados.append(r)

        if r.get("ok"):
            iv_tag    = f"IV: {iv_reais.get(base)*100:.1f}%" if iv_reais.get(base) else "IV: N/D"
            n_dets    = len(r.get("estrategias_detalhadas", []))
            n_completas = sum(1 for d in r.get("estrategias_detalhadas", []) if d.get("completa"))
            det_tag   = f"  {n_completas}/{n_dets} estratégia(s) com opções" if n_dets else ""
            print(f"  {base:<8}  {r['direcao_str']:<8}  {iv_tag}{det_tag}")

    # ── Output ────────────────────────────────────────────────────────
    ts       = datetime.now().strftime("%Y%m%d_%H%M")
    ts_human = datetime.now().strftime("%d/%m/%Y %H:%M")

    relatorio = formatar_relatorio(resultados, args.patrimonio)
    print("\n" + relatorio)

    html_path = Path(f"scanner_{ts}.html")
    html_content = gerar_html(resultados, args.patrimonio, ts_human)
    html_path.write_text(html_content, encoding="utf-8")
    print(f"\n  Relatório HTML: {html_path.resolve()}")

    if not args.no_browser:
        webbrowser.open(html_path.resolve().as_uri())
        print("  (abrindo no navegador...)")

    if args.salvar:
        txt_path  = Path(f"scanner_{ts}.txt")
        xlsx_path = Path(f"scanner_{ts}.xlsx")
        txt_path.write_text(relatorio, encoding="utf-8")
        print(f"  Relatório txt: {txt_path}")
        exportar_excel(resultados, xlsx_path)


if __name__ == "__main__":
    main()
